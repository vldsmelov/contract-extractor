from io import BytesIO
from fastapi import UploadFile
import json

from docx import Document


def _is_docx(file: UploadFile) -> bool:
    return (
        (file.filename or "").lower().endswith(".docx")
        or file.content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def _is_heading_paragraph(paragraph) -> bool:
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    return "heading" in style_name or "\u0437\u0430\u0433\u043e\u043b\u043e\u0432\u043e\u043a" in style_name


def _extract_sections_from_docx(content: bytes) -> list[str]:
    document = Document(BytesIO(content))

    sections: list[list[str]] = [[]]  # first section is the document header

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        is_heading = _is_heading_paragraph(paragraph)

        if is_heading:
            if sections[-1]:
                sections.append([])
            elif len(sections) == 1 and not sections[-1]:
                # keep the empty header, start a new section for the first heading
                sections.append([])

        sections[-1].append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    sections[-1].append(cell_text)

    return ["\n".join(chunk) for chunk in sections]


def _extract_text_from_docx(content: bytes) -> str:
    sections = _extract_sections_from_docx(content)
    non_empty_sections = [section for section in sections if section.strip()]
    return "\n\n".join(non_empty_sections) if non_empty_sections else "\n\n".join(sections)


async def read_text_and_sections_from_upload(
    file: UploadFile,
) -> tuple[str, list[str] | None]:
    content = await file.read()

    if _is_docx(file):
        try:
            sections = _extract_sections_from_docx(content)
            text = _extract_text_from_docx(content)
            return text, sections
        except Exception:
            # Fallback to decoding as plain text if DOCX parsing fails
            pass

    try:
        text = content.decode("utf-8")
    except Exception:
        text = content.decode("cp1251", errors="ignore")

    return text, None


async def read_text_from_upload(file: UploadFile) -> str:
    text, _ = await read_text_and_sections_from_upload(file)
    return text


async def read_json_from_upload(file: UploadFile):
    text = await read_text_from_upload(file)
    return json.loads(text)
